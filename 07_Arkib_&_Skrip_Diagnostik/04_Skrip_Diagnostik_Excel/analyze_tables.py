import openpyxl
import json

wb_v = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)
wb_f = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=False)

tables_detail = {}

for sname in ["MARKAH_TKRS_KAWAD", "MARKAH_TKRS_BANTU_MULA", "MARKAH_TKRS_IKATAN",
              "MARKAH_KRS_KAWAD", "MARKAH_KRS_BANTU_MULA", "MARKAH_KRS_IKATAN",
              "ANALISIS_TKRS", "ANALISIS_KRS"]:
    ws_v = wb_v[sname]
    ws_f = wb_f[sname]
    
    headers = [ws_v.cell(1, c).value for c in range(1, ws_v.max_column + 1)]
    row2 = [ws_v.cell(2, c).value for c in range(1, ws_v.max_column + 1)]
    formulas_r2 = [ws_f.cell(2, c).value for c in range(1, ws_v.max_column + 1)]
    
    tables_detail[sname] = {
        "headers": headers,
        "row2_values": row2,
        "row2_formulas": [str(f) for f in formulas_r2 if str(f).startswith('=')]
    }

with open("tables_detail.json", "w", encoding="utf-8") as f:
    json.dump(tables_detail, f, ensure_ascii=False, indent=2, default=str)

print("Tables detail dumped.")
