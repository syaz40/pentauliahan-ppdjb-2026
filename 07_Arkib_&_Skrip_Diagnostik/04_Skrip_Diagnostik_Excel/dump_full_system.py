import openpyxl
import json

wb_formulas = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=False)
wb_vals = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)

report = {}

for name in wb_formulas.sheetnames:
    ws_f = wb_formulas[name]
    ws_v = wb_vals[name]
    
    headers = [str(ws_v.cell(1, c).value or f"Col_{c}").strip() for c in range(1, ws_v.max_column + 1)]
    
    # Trim empty trailing headers
    while headers and (headers[-1].startswith("Col_") or headers[-1] == ""):
        # check if whole column is empty
        col_idx = len(headers)
        has_val = any(ws_v.cell(r, col_idx).value is not None for r in range(1, ws_v.max_row + 1))
        if not has_val:
            headers.pop()
        else:
            break
            
    rows_data = []
    formulas_used = {}
    
    for r in range(2, min(ws_v.max_row + 1, 100)):
        row_dict = {}
        has_any = False
        for c, h in enumerate(headers, 1):
            val = ws_v.cell(r, c).value
            f_val = ws_f.cell(r, c).value
            if str(f_val).startswith('='):
                formulas_used[h] = str(f_val)
            if val is not None and str(val).strip() != "":
                has_any = True
            row_dict[h] = val
        if has_any:
            rows_data.append(row_dict)
            
    report[name] = {
        "headers": headers,
        "total_rows_loaded": len(rows_data),
        "formulas": formulas_used,
        "data_preview": rows_data[:3]
    }

with open("full_system_dump.json", "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=2, default=str)

print("Full system dump created.")
