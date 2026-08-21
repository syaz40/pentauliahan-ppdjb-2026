import openpyxl
import json
from collections import Counter

wb = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)

stats = {}

# 1. PESERTA_2025 analysis
ws_p = wb["PESERTA_2025"]
p_rows = []
for r in range(2, ws_p.max_row + 1):
    pid = ws_p.cell(r, 1).value
    if pid:
        p_rows.append({
            "id": pid,
            "nama": ws_p.cell(r, 2).value,
            "sekolah": ws_p.cell(r, 3).value,
            "unit": ws_p.cell(r, 4).value,
            "ic": ws_p.cell(r, 5).value,
            "jantina": ws_p.cell(r, 6).value,
            "kategori": ws_p.cell(r, 7).value,
            "kumpulan": ws_p.cell(r, 8).value,
            "no_urutan": ws_p.cell(r, 9).value,
            "agama": ws_p.cell(r, 10).value
        })

stats["total_peserta"] = len(p_rows)
stats["by_kategori"] = dict(Counter(r["kategori"] for r in p_rows))
stats["by_unit"] = dict(Counter(r["unit"] for r in p_rows))
stats["by_jantina"] = dict(Counter(r["jantina"] for r in p_rows))
stats["by_kumpulan"] = dict(Counter(f"{r['kategori']}-{r['kumpulan']}" for r in p_rows))
stats["total_sekolah"] = len(set(r["sekolah"] for r in p_rows if r["sekolah"]))

# 2. Markah tables counts
stats["markah_counts"] = {}
for sheet_name in ["MARKAH_TKRS_KAWAD", "MARKAH_TKRS_BANTU_MULA", "MARKAH_TKRS_IKATAN",
                   "MARKAH_KRS_KAWAD", "MARKAH_KRS_BANTU_MULA", "MARKAH_KRS_IKATAN"]:
    ws = wb[sheet_name]
    count = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value:
            count += 1
    stats["markah_counts"][sheet_name] = count

# 3. Check TKRS Kawad headers
ws_tk = wb["MARKAH_TKRS_KAWAD"]
stats["tkrs_kawad_headers"] = [str(ws_tk.cell(1, c).value) for c in range(1, ws_tk.max_column + 1) if ws_tk.cell(1, c).value]

# 4. Check KRS Kawad headers
ws_kk = wb["MARKAH_KRS_KAWAD"]
stats["krs_kawad_headers"] = [str(ws_kk.cell(1, c).value) for c in range(1, ws_kk.max_column + 1) if ws_kk.cell(1, c).value]

with open("peserta_and_markah_stats.json", "w", encoding="utf-8") as f:
    json.dump(stats, f, ensure_ascii=False, indent=2)

print("Stats processed successfully.")
