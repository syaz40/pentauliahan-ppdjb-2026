import openpyxl

wb = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)

for sheet_name in ["MARKAH_TKRS_KAWAD", "MARKAH_KRS_KAWAD"]:
    ws = wb[sheet_name]
    print(f"=== {sheet_name} Columns ===")
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        # Letter
        col_letter = openpyxl.utils.get_column_letter(c)
        if h or c <= 60:
            print(f"Col {c} ({col_letter}): {h}")

