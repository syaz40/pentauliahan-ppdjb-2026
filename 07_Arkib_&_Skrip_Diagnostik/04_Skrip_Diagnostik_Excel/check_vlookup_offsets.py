import openpyxl

wb = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)

sheets = [
    "MARKAH_TKRS_KAWAD",
    "MARKAH_TKRS_BANTU_MULA",
    "MARKAH_TKRS_IKATAN",
    "MARKAH_KRS_KAWAD",
    "MARKAH_KRS_BANTU_MULA",
    "MARKAH_KRS_IKATAN"
]

print("=== CHECKING COLUMNS FROM COLUMN B (Index 1 = Col B) ===")
for s in sheets:
    ws = wb[s]
    print(f"\n--- Sheet: {s} (Total max col: {ws.max_column}) ---")
    for c in range(2, ws.max_column + 1):
        idx_from_B = c - 1  # Col B is index 1
        col_letter = openpyxl.utils.get_column_letter(c)
        h = ws.cell(1, c).value
        if h:
            print(f"  Index {idx_from_B} (Col {col_letter}): {h}")

