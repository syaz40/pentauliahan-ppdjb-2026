import openpyxl
from openpyxl.utils import get_column_letter

# Load the current workbook
wb_source = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=False)
wb_vals = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)

# Create a clean revised workbook
wb_new = openpyxl.Workbook()
# remove default sheet
wb_new.remove(wb_new.active)

# Essential sheets to retain and fix
valid_sheets = [
    "PESERTA_2025",
    "JENIS_PENTAULIAHAN",
    "MARKAH_TKRS_KAWAD",
    "MARKAH_TKRS_BANTU_MULA",
    "MARKAH_TKRS_IKATAN",
    "MARKAH_KRS_KAWAD",
    "MARKAH_KRS_BANTU_MULA",
    "MARKAH_KRS_IKATAN",
    "ANALISIS_TKRS",
    "ANALISIS_KRS",
    "PESERTA TIDAK HADIR"
]

for sname in valid_sheets:
    if sname not in wb_source.sheetnames:
        continue
        
    ws_src_f = wb_source[sname]
    ws_src_v = wb_vals[sname]
    
    ws_new = wb_new.create_sheet(title=sname)
    
    if sname == "ANALISIS_TKRS":
        # Create single clean header row
        headers = [
            "BIL", "ID_PESERTA", "NAMA", "SEKOLAH", "JANTINA", "AGAMA",
            "MARKAH_KAWAD_KAKI", "MARKAH_BANTU_MULA", "MARKAH_IKATAN",
            "JUMLAH_MARKAH", "STATUS_TKRS", "STATUS_PENUH"
        ]
        ws_new.append(headers)
        
        # Populate data starting from source row 3 (skip duplicate subheaders in row 2)
        target_row = 2
        for r in range(3, ws_src_v.max_row + 1):
            pid = ws_src_v.cell(r, 2).value
            if not pid:
                continue
            bil = ws_src_v.cell(r, 1).value or (target_row - 1)
            nama = ws_src_v.cell(r, 3).value
            sekolah = ws_src_v.cell(r, 4).value
            jantina = ws_src_v.cell(r, 5).value
            agama = ws_src_v.cell(r, 6).value
            
            # FIXED FORMULAS FOR TKRS
            # TKRS Kawad is Col AO (Col 41), offset from Col B is index 40. Full score = 95
            f_kawad = f'=IFERROR(VLOOKUP($B{target_row}, MARKAH_TKRS_KAWAD!$B:$AP, 40, FALSE)/95, 0)'
            # TKRS Bantu Mula is Col J (Col 10), offset from Col B is index 9. Full score = 20
            f_bantu = f'=IFERROR(VLOOKUP($B{target_row}, MARKAH_TKRS_BANTU_MULA!$B:$K, 9, FALSE)/20, 0)'
            # TKRS Ikatan is Col L (Col 12), offset from Col B is index 11. Full score = 25
            f_ikatan = f'=IFERROR(VLOOKUP($B{target_row}, MARKAH_TKRS_IKATAN!$B:$M, 11, FALSE)/25, 0)'
            # Jumlah Markah Wajaran: Kawad 60%, Bantu Mula 20%, Ikatan 20%
            f_jumlah = f'=IFERROR((G{target_row}*0.6) + (H{target_row}*0.2) + (I{target_row}*0.2), 0)'
            # Status Kod 3 Huruf: LLL / LGG / GGG
            f_status_kod = f'=IF(G{target_row}>=0.5,"L","G") & IF(((H{target_row}+I{target_row})/2)>=0.5,"L","G") & IF(J{target_row}>=0.6,"L","G")'
            # Status Teks Mesra Pengguna
            f_status_penuh = f'=IFS(J{target_row}=0, "BELUM LENGKAP", K{target_row}="LLL", "LULUS PENTAULIAHAN", J{target_row}>=0.5, "LULUS BERSYARAT", TRUE, "GAGAL")'
            
            ws_new.append([
                bil, pid, nama, sekolah, jantina, agama,
                f_kawad, f_bantu, f_ikatan, f_jumlah, f_status_kod, f_status_penuh
            ])
            target_row += 1
            
    elif sname == "ANALISIS_KRS":
        # Create single clean header row
        headers = [
            "BIL", "ID_PESERTA", "NAMA", "SEKOLAH", "JANTINA", "AGAMA",
            "MARKAH_KAWAD_KAKI", "MARKAH_BANTU_MULA", "MARKAH_IKATAN",
            "JUMLAH_MARKAH", "STATUS_KRS", "STATUS_PENUH"
        ]
        ws_new.append(headers)
        
        target_row = 2
        for r in range(3, ws_src_v.max_row + 1):
            pid = ws_src_v.cell(r, 2).value
            if not pid:
                continue
            bil = ws_src_v.cell(r, 1).value or (target_row - 1)
            nama = ws_src_v.cell(r, 3).value
            sekolah = ws_src_v.cell(r, 4).value
            jantina = ws_src_v.cell(r, 5).value
            agama = ws_src_v.cell(r, 6).value
            
            # FIXED FORMULAS FOR KRS
            # KRS Kawad is Col BC (Col 55), offset from Col B is index 54. Full score = 137
            f_kawad = f'=IFERROR(VLOOKUP($B{target_row}, MARKAH_KRS_KAWAD!$B:$BD, 54, FALSE)/137, 0)'
            # KRS Bantu Mula is Col K (Col 11), offset from Col B is index 10. Full score = 20
            f_bantu = f'=IFERROR(VLOOKUP($B{target_row}, MARKAH_KRS_BANTU_MULA!$B:$K, 10, FALSE)/20, 0)'
            # KRS Ikatan is Col M (Col 13), offset from Col B is index 12. Full score = 25
            f_ikatan = f'=IFERROR(VLOOKUP($B{target_row}, MARKAH_KRS_IKATAN!$B:$M, 12, FALSE)/25, 0)'
            # Jumlah Markah Wajaran
            f_jumlah = f'=IFERROR((G{target_row}*0.6) + (H{target_row}*0.2) + (I{target_row}*0.2), 0)'
            # Status Kod
            f_status_kod = f'=IF(G{target_row}>=0.5,"L","G") & IF(((H{target_row}+I{target_row})/2)>=0.5,"L","G") & IF(J{target_row}>=0.6,"L","G")'
            # Status Teks Mesra Pengguna
            f_status_penuh = f'=IFS(J{target_row}=0, "BELUM LENGKAP", K{target_row}="LLL", "LULUS PENTAULIAHAN", J{target_row}>=0.5, "LULUS BERSYARAT", TRUE, "GAGAL")'
            
            ws_new.append([
                bil, pid, nama, sekolah, jantina, agama,
                f_kawad, f_bantu, f_ikatan, f_jumlah, f_status_kod, f_status_penuh
            ])
            target_row += 1
            
    else:
        # Copy original table data cleanly without trailing empty columns
        max_c = ws_src_v.max_column
        # find last column with header
        actual_headers = []
        for c in range(1, max_c + 1):
            h_val = ws_src_v.cell(1, c).value
            if h_val:
                actual_headers.append((c, str(h_val).strip()))
                
        # Write headers
        ws_new.append([h[1] for h in actual_headers])
        
        # Write data rows
        for r in range(2, ws_src_v.max_row + 1):
            row_data = []
            has_val = False
            for col_idx, _ in actual_headers:
                f_val = ws_src_f.cell(r, col_idx).value
                v_val = ws_src_v.cell(r, col_idx).value
                if str(f_val).startswith('='):
                    row_data.append(str(f_val))
                else:
                    row_data.append(v_val)
                if v_val is not None:
                    has_val = True
            if has_val:
                ws_new.append(row_data)

wb_new.save("PENTAULIAHAN_BACKEND_REVISED.xlsx")
print("PENTAULIAHAN_BACKEND_REVISED.xlsx generated successfully!")
