import openpyxl

wb_f = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=False)
wb_v = openpyxl.load_workbook("backend_pentauliahan.xlsx", data_only=True)

for sheet in ["ANALISIS_TKRS", "ANALISIS_KRS"]:
    ws_f = wb_f[sheet]
    ws_v = wb_v[sheet]
    print(f"=== {sheet} First 5 Rows ===")
    for r in range(1, 6):
        row_vals = []
        for c in range(1, 12):
            raw = ws_f.cell(r, c).value
            val = ws_v.cell(r, c).value
            if str(raw).startswith('='):
                row_vals.append(f"[{raw} => {val}]")
            else:
                row_vals.append(str(val))
        print(f"Row {r}: " + " | ".join(row_vals))

