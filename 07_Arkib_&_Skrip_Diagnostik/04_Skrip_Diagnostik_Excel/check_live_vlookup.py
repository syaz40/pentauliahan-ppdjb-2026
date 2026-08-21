import openpyxl

wb = openpyxl.load_workbook('live_google_sheet.xlsx', data_only=False)

def check_sheet_columns(sheet_name):
    ws = wb[sheet_name]
    headers = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        headers.append((c, openpyxl.utils.get_column_letter(c), v))
    return headers

print("=== MARKAH_TKRS_KAWAD COLUMNS ===")
for col_idx, letter, name in check_sheet_columns('MARKAH_TKRS_KAWAD'):
    if name is not None or col_idx in [1, 2, 40, 41, 42, 43, 44, 45, 54, 55, 56]:
        print(f"Col {col_idx} ({letter}): {name}")

print("\n=== MARKAH_KRS_KAWAD COLUMNS ===")
for col_idx, letter, name in check_sheet_columns('MARKAH_KRS_KAWAD'):
    if name is not None or col_idx in [1, 2, 50, 51, 52, 53, 54, 55, 56, 57]:
        print(f"Col {col_idx} ({letter}): {name}")

print("\n=== ANALISIS_TKRS ROW 1 & ROW 2 ===")
ws_a = wb['ANALISIS_TKRS']
for c in range(1, ws_a.max_column + 1):
    r1 = ws_a.cell(row=1, column=c).value
    r2 = ws_a.cell(row=2, column=c).value
    r3 = ws_a.cell(row=3, column=c).value
    print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): R1='{r1}' | R2='{r2}' | R3 formula='{r3}'")

print("\n=== ANALISIS_KRS ROW 1 & ROW 2 ===")
ws_k = wb['ANALISIS_KRS']
for c in range(1, ws_k.max_column + 1):
    r1 = ws_k.cell(row=1, column=c).value
    r2 = ws_k.cell(row=2, column=c).value
    r3 = ws_k.cell(row=3, column=c).value
    print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): R1='{r1}' | R2='{r2}' | R3 formula='{r3}'")
