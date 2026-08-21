import openpyxl
import json

wb = openpyxl.load_workbook('live_google_sheet.xlsx', data_only=False)
wb_val = openpyxl.load_workbook('live_google_sheet.xlsx', data_only=True)

print('Sheet names in Live Google Sheet:', wb.sheetnames)

sheet_info = {}
for name in wb.sheetnames:
    ws = wb[name]
    ws_v = wb_val[name]
    
    headers = [ws_v.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    while headers and headers[-1] is None:
        headers.pop()
        
    non_empty_rows = 0
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value is not None for c in range(1, len(headers)+1)):
            non_empty_rows += 1
            
    formulas = {}
    for c_idx, h in enumerate(headers, 1):
        for r in range(2, min(ws.max_row + 1, 10)):
            val = ws.cell(row=r, column=c_idx).value
            if str(val).startswith('='):
                formulas[str(h)] = str(val)
                break
                
    sheet_info[name] = {
        'total_data_rows': non_empty_rows,
        'total_columns': len(headers),
        'headers': headers,
        'formulas': formulas
    }

print('\n=== SUMMARY OF SHEETS ===')
for name, info in sheet_info.items():
    r_count = info['total_data_rows']
    c_count = info['total_columns']
    print(f"Sheet: [{name}] | Rows: {r_count} | Cols: {c_count}")
    if info['formulas']:
        print(f"   Formulas: {info['formulas']}")

with open('live_sheet_summary.json', 'w', encoding='utf-8') as f:
    json.dump(sheet_info, f, ensure_ascii=False, indent=2, default=str)

print('\nSaved to live_sheet_summary.json successfully!')
