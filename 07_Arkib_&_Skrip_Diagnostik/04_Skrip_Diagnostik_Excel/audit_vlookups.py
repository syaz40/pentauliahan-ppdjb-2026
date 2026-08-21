import openpyxl

wb = openpyxl.load_workbook('live_google_sheet.xlsx', data_only=False)
wb_val = openpyxl.load_workbook('live_google_sheet.xlsx', data_only=True)

def analyze_vlookup_targets(sheet_name, lookup_sheet, vlookup_range, col_index):
    ws_src = wb[lookup_sheet]
    ws_src_v = wb_val[lookup_sheet]
    
    # parse range like B:BD or $B:$L
    start_col_letter = vlookup_range.split(':')[0].replace('$', '')
    start_col_num = openpyxl.utils.column_index_from_string(start_col_letter)
    target_col_num = start_col_num + col_index - 1
    target_col_letter = openpyxl.utils.get_column_letter(target_col_num)
    
    header_val = ws_src_v.cell(row=1, column=target_col_num).value
    sample_val = ws_src_v.cell(row=2, column=target_col_num).value
    
    return {
        'start_col': f"{start_col_letter} ({start_col_num})",
        'col_index': col_index,
        'target_col': f"{target_col_letter} ({target_col_num})",
        'header_at_target': header_val,
        'sample_val_row2': sample_val
    }

print("=== VLOOKUP AUDIT IN ANALISIS_TKRS ===")
print("1. KAWAD:", analyze_vlookup_targets('ANALISIS_TKRS', 'MARKAH_TKRS_KAWAD', 'B:BD', 54))
print("2. BANTU MULA:", analyze_vlookup_targets('ANALISIS_TKRS', 'MARKAH_TKRS_BANTU_MULA', 'B:L', 11))
print("3. IKATAN:", analyze_vlookup_targets('ANALISIS_TKRS', 'MARKAH_TKRS_IKATAN', 'B:N', 12))

print("\n=== VLOOKUP AUDIT IN ANALISIS_KRS ===")
print("1. KAWAD:", analyze_vlookup_targets('ANALISIS_KRS', 'MARKAH_KRS_KAWAD', 'B:BD', 54))
print("2. BANTU MULA:", analyze_vlookup_targets('ANALISIS_KRS', 'MARKAH_KRS_BANTU_MULA', 'B:L', 11))
print("3. IKATAN:", analyze_vlookup_targets('ANALISIS_KRS', 'MARKAH_KRS_IKATAN', 'B:N', 12))
